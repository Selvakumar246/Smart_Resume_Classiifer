import io
import re
from pathlib import Path
from fastapi import HTTPException, UploadFile
from docx import Document

# Safe import for PyMuPDF
try:
    import pymupdf
except ImportError:
    try:
        import fitz as pymupdf
    except ImportError:
        pymupdf = None

from app.core.config import get_settings

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}

async def extract_resume_text(file: UploadFile) -> str:
    settings = get_settings()
    filename = file.filename or "resume"
    extension = Path(filename).suffix.lower()
    
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail="Upload a PDF, DOCX, or TXT resume.")

    data = await file.read()
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(status_code=413, detail=f"File exceeds {settings.max_upload_mb} MB.")

    try:
        if extension == ".pdf":
            if pymupdf is None:
                raise ImportError("PyMuPDF (pymupdf/fitz) is not installed.")
            # Open PDF with PyMuPDF
            doc = pymupdf.open(stream=io.BytesIO(data), filetype="pdf")
            text = "\n".join(page.get_text() or "" for page in doc)
            doc.close()
        elif extension == ".docx":
            document = Document(io.BytesIO(data))
            text_list = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    text_list.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(text_list)
        else:
            text = data.decode("utf-8", errors="ignore")
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail=f"The resume could not be parsed. Try exporting it again. Error: {str(exc)}"
        ) from exc

    # Clean and normalize text
    # Replace non-breaking spaces and other weird whitespaces with standard space
    cleaned = re.sub(r"[\xa0\t]+", " ", text)
    # Deduplicate empty lines but preserve paragraphs (max 2 consecutive newlines)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    
    # Check length
    if len(cleaned) < 80:
        raise HTTPException(status_code=422, detail="Not enough readable text was found in the file.")
        
    return cleaned
